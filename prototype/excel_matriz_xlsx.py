# -*- coding: utf-8 -*-
"""
excel_matriz_xlsx.py — a matriz que está NA TELA, como .xlsx de verdade.
=========================================================================
[2026-09-25, pedido do usuário: "o formato do excel não é XLSX, pode corrigir
e validar?"] O botão "⬇ Baixar Excel" gerava o arquivo no próprio navegador,
em SpreadsheetML 2003 (XML com extensão .xls) — formato de 2003 que o Excel
abre, mas que não é Open XML: o Excel avisa que o conteúdo não bate com a
extensão, e ferramenta que lê xlsx de verdade (pandas, Google Sheets, viewer
de e-mail) recusa. Agora quem escreve é este módulo, com `openpyxl` — a MESMA
biblioteca que o relatório do CLI já usava (`excel_report.py`), então não
entrou dependência nova (CLAUDE.md §13).

Divisão de responsabilidade, de propósito:
  - a TELA decide O QUE entra (linhas filtradas, na ordem em tela, com as
    anotações inclusive as ainda não salvas, os alertas dia a dia, o
    comentário vigente de cada dia) e manda isso pronto em JSON —
    static/js/controle_cargas/exportar.js;
  - este módulo só FORMATA (cor, contorno, congelamento, largura). Nada de
    regra de negócio aqui: ele não sabe o que é "pauta", só recebe um
    booleano por célula.
Isso mantém o "o que você vê é o que você baixa" que o gerador no navegador
tinha — o servidor não conhece os filtros nem a ordenação da tela.

A paleta vem de `excel_report._PREENCHIMENTO_XLSX`: os dois geradores de
Excel do app passam a ler do MESMO lugar (antes era um dict em Python e outro
igualzinho em JS, que só um comentário mantinha em sincronia).
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from excel_report import _PREENCHIMENTO_XLSX

# Contorno das células com badge "Pauta do dia" — mesmo fúcsia do badge na
# tela (--overlay-pauta) [2026-09-24, pedido do usuário].
COR_CONTORNO_PAUTA = "C026D3"

# Colunas de identidade, na ordem em que aparecem (as 4 primeiras ficam
# congeladas com o cabeçalho).
COLUNAS_IDENTIDADE = ["Company", "Carteira", "WalletID", "Instituição",
                      "Modelo de Carga", "Periodicidade", "Defasagem"]

# Fundo/letra do comentário vigente, aplicados nas 4 primeiras colunas —
# mesmos hex de XML_BG_COMENTARIO/XML_FG_COMENTARIO (exportar.js).
_PREENCHIMENTO_COMENTARIO = {
    "red": ("FEE2E2", "B91C1C"),
    "yellow": ("FEF3C7", "92400E"),
    "green": ("DCFCE7", "15803D"),
}

_LARGURAS = {"Company": 22, "Carteira": 38, "WalletID": 26, "Instituição": 18,
             "Modelo de Carga": 16, "Periodicidade": 12, "Defasagem": 22}
_LARGURA_DIA = 7
_LARGURA_AUDITORIA = 42


def _borda_pauta():
    """Contexto:
    Borda fúcsia grossa dos 4 lados, usada nas células com badge Pauta.
    Chamada 1x por escrita (o objeto é reaproveitado em todas as células).
    Retorna um `Border` do openpyxl.

    Pseudocódigo:
      1. Monta um lado "grosso" na cor da pauta.
      2. Aplica o mesmo lado nas 4 direções.
    """
    lado = Side(style="thick", color=COR_CONTORNO_PAUTA)
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def _escrever_cabecalho(aba, payload):
    """Contexto:
    Escreve as 2 primeiras linhas: a de geração (com a legenda do contorno) e
    a de cabeçalho das colunas. Chamada por montar_workbook_matriz(). Retorna
    a lista de títulos escritos (quem chama usa pra largura/autofiltro).

    Pseudocódigo:
      1. Linha 1: quando foi gerado + o que significa o contorno fúcsia.
      2. Linha 2: identidade + 1 coluna por dia da janela + as colunas de
         auditoria, que cobrem a janela inteira (rótulo com o range).
      3. Negrito e fundo cinza no cabeçalho.
    """
    janela = payload["janela"]
    rotulo_janela = payload.get("rotuloJanela") or f"{janela[0]}–{janela[-1]}"
    referencia = payload.get("referenceDate") or ""

    aba.append([f"Relatório gerado em: {payload.get('geradoEm', '')} — células com contorno "
                f"fúcsia = Pauta do dia (dia exato da Defasagem); as colunas de alerta cobrem "
                f"toda a janela, dia a dia"])

    titulos = (COLUNAS_IDENTIDADE + list(janela) + [
        f"Responsável ({referencia})", f"Comentário sobre atuação ({referencia})",
        f"Δ Rent bp ({rotulo_janela})", f"Fora de sequência ({rotulo_janela})",
        f"Issues ({rotulo_janela})",
        f"Alertas do Grid — Rent/Atraso/Pauta/Sequência/Issue ({rotulo_janela})",
        f"Comentário ({rotulo_janela})"])
    aba.append(titulos)

    preenchimento = PatternFill("solid", fgColor="EEF0EC")
    for coluna in range(1, len(titulos) + 1):
        celula = aba.cell(row=2, column=coluna)
        celula.font = Font(bold=True)
        celula.fill = preenchimento
    return titulos


def _escrever_linha(aba, numero_linha, linha, janela, borda_pauta):
    """Contexto:
    Escreve 1 carteira: identidade (tingida pela severidade do comentário
    vigente), as células de dia (fundo = estágio, contorno = pauta) e as
    colunas de auditoria. Chamada por montar_workbook_matriz() por linha.
    Não retorna nada.

    Pseudocódigo:
      1. Identidade; nas 4 primeiras colunas, aplica a cor do comentário
         vigente quando a tela mandou uma severidade.
      2. 1 célula por dia: sigla no centro, fundo/letra do estágio e, quando
         o dia é pauta, o contorno fúcsia.
      3. As 7 colunas de auditoria, como texto.
    """
    valores_identidade = [linha.get(c, "") for c in
                          ("company", "carteira", "walletId", "instituicao",
                           "modeloCarga", "periodicidade", "defasagem")]
    for indice, valor in enumerate(valores_identidade, start=1):
        aba.cell(row=numero_linha, column=indice, value=valor)

    severidade = linha.get("severidadeComentario")
    if severidade in _PREENCHIMENTO_COMENTARIO:
        fundo, letra = _PREENCHIMENTO_COMENTARIO[severidade]
        for indice in range(1, 5):
            celula = aba.cell(row=numero_linha, column=indice)
            celula.fill = PatternFill("solid", fgColor=fundo)
            celula.font = Font(color=letra)

    primeira_coluna_dia = len(COLUNAS_IDENTIDADE) + 1
    dias = linha.get("dias") or []
    for indice, dia in enumerate(dias[:len(janela)]):
        celula = aba.cell(row=numero_linha, column=primeira_coluna_dia + indice,
                          value=dia.get("letra", ""))
        celula.alignment = Alignment(horizontal="center")
        fundo, letra = _PREENCHIMENTO_XLSX.get(dia.get("estado"), _PREENCHIMENTO_XLSX["notcov"])
        celula.fill = PatternFill("solid", fgColor=fundo)
        celula.font = Font(color=letra, bold=True)
        if dia.get("pauta"):
            celula.border = borda_pauta

    primeira_coluna_auditoria = primeira_coluna_dia + len(janela)
    for indice, chave in enumerate(("responsavel", "comentarioAtuacao", "divergencia",
                                    "sequencia", "issues", "alertas", "comentarios")):
        aba.cell(row=numero_linha, column=primeira_coluna_auditoria + indice,
                 value=linha.get(chave, ""))


def _ajustar_colunas(aba, titulos, total_janela):
    """Contexto:
    Larguras, congelamento e autofiltro — o "chrome" da planilha. Chamada 1x
    no fim de montar_workbook_matriz(). Não retorna nada.

    Pseudocódigo:
      1. Largura fixa por tipo de coluna (identidade / dia / auditoria).
      2. Congela as 4 primeiras colunas e as 2 primeiras linhas (mesmo
         congelamento que o gerador antigo fazia).
      3. Autofiltro na linha de cabeçalho.
    """
    for indice, titulo in enumerate(titulos, start=1):
        if indice <= len(COLUNAS_IDENTIDADE):
            largura = _LARGURAS.get(titulo, 18)
        elif indice <= len(COLUNAS_IDENTIDADE) + total_janela:
            largura = _LARGURA_DIA
        else:
            largura = _LARGURA_AUDITORIA
        aba.column_dimensions[get_column_letter(indice)].width = largura

    aba.freeze_panes = "E3"
    aba.auto_filter.ref = f"A2:{get_column_letter(len(titulos))}{max(aba.max_row, 2)}"


def montar_workbook_matriz(payload):
    """Contexto:
    Monta o .xlsx da aba Matriz a partir do que a TELA mandou (já filtrado e
    ordenado) — ponto de entrada deste módulo, chamado pela rota
    `POST /api/exportar-excel` (app.py) [2026-09-25, pedido do usuário: "o
    formato do excel não é XLSX"]. Retorna um BytesIO posicionado no começo,
    pronto pra virar resposta HTTP.

    Pseudocódigo:
      1. Cria a pasta de trabalho com a aba "Matriz".
      2. Escreve as 2 linhas de cabeçalho.
      3. Escreve 1 linha por carteira recebida.
      4. Ajusta largura/congelamento/autofiltro.
      5. Salva em memória e devolve o buffer.
    """
    pasta = Workbook()
    aba = pasta.active
    aba.title = "Matriz"

    titulos = _escrever_cabecalho(aba, payload)
    borda_pauta = _borda_pauta()
    for numero, linha in enumerate(payload.get("linhas") or [], start=3):
        _escrever_linha(aba, numero, linha, payload["janela"], borda_pauta)

    _ajustar_colunas(aba, titulos, len(payload["janela"]))

    buffer = BytesIO()
    pasta.save(buffer)
    buffer.seek(0)
    return buffer
